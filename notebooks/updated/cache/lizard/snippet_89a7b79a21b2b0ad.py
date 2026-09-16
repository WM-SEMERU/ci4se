def write_sequences_to_xlsx(path, seqs):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for row, id in enumerate(seqs, 1):
        ws.cell(row, 1).value = id
        ws.cell(row, 2).value = seqs[id]
    wb.save(path)